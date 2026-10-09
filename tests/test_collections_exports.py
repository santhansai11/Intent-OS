"""Exercise real local collections/exports and provider boundaries using synthetic files."""
import hashlib
import io
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile
import xml.etree.ElementTree as ET

from intent_kernel.config import Config
from intent_kernel.runtime import Runtime
from planner.router import PlannerRouter
from services.capabilities import CapabilityRegistry
from services.documents import FORMATS
from services.gemini import Gemini, AIKeys
from services.connections import ConnectionError


class LocalWorkflows(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.config = Config(Path(self.temp.name) / 'runtime')
        self.registry = CapabilityRegistry(self.config, None)
        self.source = self.config.workspace / 'originals'
        self.source.mkdir()
        (self.source / 'one').mkdir()
        (self.source / 'two').mkdir()
        (self.source / 'one' / 'notes.txt').write_text('Cloud computing meeting: elasticity and shared responsibility.')
        (self.source / 'two' / 'notes.txt').write_text('Cloud computing cost review. Capacity planning.')
        (self.source / 'holiday.txt').write_text('Beach and sunshine.')
        self.task = {'roots': [str(self.config.workspace)], 'scopes': ['files'], 'access_mode': 'selected'}

    def tearDown(self):
        self.temp.cleanup()

    def call(self, tool, **arguments):
        return self.registry.execute({'tool': tool, 'arguments': arguments}, self.task)

    def manifest(self):
        return {str(p.relative_to(self.source)): hashlib.sha256(p.read_bytes()).hexdigest() for p in self.source.rglob('*') if p.is_file()}

    def test_collect_content_matches_collision_safe_copies_and_archive(self):
        before = self.manifest()
        result = self.call('files_collect', query='cloud computing', destination='meeting-pack')
        self.assertEqual(result['copied'], 2)
        self.assertFalse(result['truncated'])
        self.assertEqual(len({Path(f['copy']).name for f in result['files']}), 2)
        for item in result['files']:
            self.assertEqual(Path(item['path']).read_bytes(), Path(item['copy']).read_bytes())
            self.assertEqual(item['sha256'], hashlib.sha256(Path(item['copy']).read_bytes()).hexdigest())
        self.assertEqual(before, self.manifest())
        with zipfile.ZipFile(result['artifact']) as bundle:
            self.assertEqual(len(bundle.namelist()), 3)
            self.assertEqual(json.loads(bundle.read('manifest.json'))['status'], 'complete')

    def test_cache_reuses_extraction_and_invalidates_changed_files(self):
        first = self.call('files_find_related', query='cloud')
        with patch.object(self.registry, 'read', side_effect=AssertionError('Cache should be used')):
            second = self.call('files_find_related', query='cloud')
        self.assertEqual(first['matched'], second['matched'])
        self.assertEqual(second['cache_hits'], 3)
        (self.source / 'holiday.txt').write_text('Cloud service migration discussion with fresh content')
        self.assertEqual(self.call('files_find_related', query='cloud')['matched'], 3)

    def test_topic_phrases_do_not_match_generic_words_or_substrings(self):
        fixtures = {
            'Operating_Systems.txt': 'Scheduling and virtual memory.',
            'OS-notes.txt': 'Process scheduling.',
            'database.txt': 'Database system architecture and data storage.',
            'aws.txt': 'AWS storage system. Store that data at a location.',
            'random.txt': 'The workshop is operating at capacity.',
            'claws.txt': 'Claws and jaws are animal features.',
            'actual-notes.txt': 'An operating system manages processes.',
        }
        for name, content in fixtures.items():
            (self.source / name).write_text(content)
        found = self.call('files_find_related', query='gather all operating system related files')
        self.assertEqual(found['effective_topics'], ['operating system'])
        self.assertEqual({m['name'] for m in found['matches']}, {'Operating_Systems.txt', 'actual-notes.txt'})
        aliases = self.call('files_find_related', query='operating system, OS')
        self.assertEqual({m['name'] for m in aliases['matches']}, {'Operating_Systems.txt', 'OS-notes.txt', 'actual-notes.txt'})
        aws = self.call('files_find_related', query='aws')
        self.assertEqual([m['name'] for m in aws['matches']], ['aws.txt'])

    def test_instruction_with_destination_is_rejected_before_creating_task(self):
        with patch('services.connections.CredentialStore.get', return_value={}):
            kernel = Runtime(self.config)
        try:
            with self.assertRaisesRegex(ValueError, 'New collection folder'):
                kernel.submit('Gather files', auto_outputs=True, workflow={'type': 'collect',
                    'query': r'gather all operating system related files and store that at D:\meeting',
                    'destination': 'wrong-auto-folder'})
            self.assertEqual(len(kernel.tasks), 0)
            self.assertFalse((self.config.workspace / 'wrong-auto-folder').exists())
        finally:
            kernel.close()

    def test_phrase_preserves_internal_function_words(self):
        (self.source / 'iot.txt').write_text('The Internet of Things connects physical devices.')
        found = self.call('files_find_related', query='internet of things')
        self.assertEqual([m['name'] for m in found['matches']], ['iot.txt'])

    def test_existing_folder_never_changed_and_empty_match_has_no_folder(self):
        before = self.manifest()
        with self.assertRaises(ValueError):
            self.call('files_collect', query='cloud', destination=str(self.source))
        self.assertEqual(before, self.manifest())
        result = self.call('files_collect', query='volcanology', destination='empty-output')
        self.assertEqual(result['status'], 'no_matches')
        self.assertFalse((self.config.workspace / 'empty-output').exists())

    def test_scan_batches_continue_past_46_files_without_partial_completion(self):
        self.config.values['scan_batch_items'] = 10
        for index in range(70):
            (self.source / f'sample-{index:03}.txt').write_text('ordinary unrelated data')
        (self.source / 'z-last-dbms.txt').write_text('Database management systems notes.')
        self.task['pending'] = {'id': 'batch-regression'}
        result = self.call('files_find_related', query='DBMS')
        self.assertEqual(result['status'], 'scanning')
        for _ in range(100):
            result = self.call('files_find_related', query='DBMS')
            if result['status'] != 'scanning':
                break
        self.assertEqual(result['phase'], 'complete')
        self.assertFalse(result['truncated'])
        self.assertEqual(result['scanned'], 74)
        self.assertEqual([m['name'] for m in result['matches']], ['z-last-dbms.txt'])

    def test_scan_checkpoint_survives_restart_and_copies_only_after_search(self):
        self.config.values['scan_batch_items'] = 1
        with patch('services.connections.CredentialStore.get', return_value={}):
            kernel = Runtime(self.config)
        try:
            task = kernel.submit('Gather cloud', auto_outputs=True,
                workflow={'type': 'collect', 'query': 'cloud', 'destination': 'resumed-pack'})
            kernel.tick()
            live = kernel.tasks[task['id']]
            self.assertEqual(live['state'], 'READY')
            self.assertEqual(live['steps'], [])
            self.assertFalse((self.config.workspace / 'resumed-pack').exists())
            # Simulate a crash during the next read-only scanning batch.
            live['state'] = 'RUNNING'
            kernel._save(live, 'dispatch', 'Scanning interrupted')
            kernel.close()
            with patch('services.connections.CredentialStore.get', return_value={}):
                kernel = Runtime(self.config)
            self.assertEqual(kernel.tasks[task['id']]['state'], 'READY')
            kernel.control(task['id'], 'pause')
            self.assertEqual(kernel.tasks[task['id']]['state'], 'PAUSED')
            kernel.control(task['id'], 'resume')
            for _ in range(50):
                kernel.tick()
                if kernel.tasks[task['id']]['state'] == 'COMPLETED':
                    break
            final = kernel.tasks[task['id']]
            self.assertEqual(final['state'], 'COMPLETED')
            self.assertEqual(final['turns'], 0)
            self.assertEqual(len(final['steps']), 1)
            self.assertEqual(final['steps'][0]['result']['copied'], 2)
            self.assertNotIn('scan_progress', final)
        finally:
            kernel.close()

    def test_all_filenames_discovered_before_document_extraction(self):
        self.config.values['scan_batch_items'] = 1
        self.task['pending'] = {'id': 'filename-first'}
        with patch.object(self.registry, 'read', wraps=self.registry.read) as read:
            for _ in range(20):
                result = self.call('files_find_related', query='cloud')
                if result['phase'] != 'discover':
                    break
                read.assert_not_called()
            self.assertEqual(result['scanned'], 3)

    def test_time_budget_yields_without_finishing_or_losing_scan_position(self):
        from itertools import count
        self.task['pending'] = {'id': 'slow-scan'}
        with patch('services.file_scan.time.monotonic', side_effect=count(0, 30)):
            first = self.call('files_find_related', query='cloud')
            self.assertEqual(first['status'], 'scanning')
            self.assertFalse(first['truncated'])
            for _ in range(30):
                result = self.call('files_find_related', query='cloud')
                if result['status'] != 'scanning':
                    break
        self.assertEqual(result['phase'], 'complete')
        self.assertEqual(result['matched'], 2)
        self.assertGreater(result['seconds'], 20)

    def test_prior_collections_are_excluded_from_future_scans(self):
        self.call('files_collect', query='cloud', destination='old-pack')
        result = self.call('files_find_related', query='cloud')
        self.assertEqual(result['matched'], 2)
        self.assertEqual(result['excluded_generated'], 1)

    def test_output_grant_does_not_expand_selected_search_folders(self):
        (self.config.workspace / 'old-report.txt').write_text('cloud computing')
        self.task['search_roots'] = [str(self.source)]
        result = self.call('files_collect', query='cloud', destination='source-only-pack')
        self.assertEqual(result['copied'], 2)
        self.assertTrue(all(Path(f['path']).is_relative_to(self.source) for f in result['files']))

    def test_custom_output_folder_outside_sources_is_granted_without_searching_it(self):
        output = Path(self.temp.name) / 'external-meeting'
        output.mkdir()
        sentinel = output / 'existing-cloud.txt'
        sentinel.write_text('cloud computing unrelated existing output')
        self.config.update({'roots': [str(self.source)]})
        with patch('services.connections.CredentialStore.get', return_value={}):
            kernel = Runtime(self.config)
        try:
            task = kernel.submit('Gather cloud files', roots=[str(self.source)], output_folder=str(output),
                auto_outputs=True, workflow={'type': 'collect', 'query': 'cloud computing', 'destination': 'new-pack'})
            self.assertEqual(task['search_roots'], [str(self.source)])
            self.assertIn(str(output), task['roots'])
            self.assertNotIn(str(output.parent), task['roots'])
            kernel.tick()
            result = kernel.tasks[task['id']]
            self.assertEqual(result['state'], 'COMPLETED')
            self.assertEqual(result['steps'][0]['result']['copied'], 2)
            self.assertEqual(result['steps'][0]['result']['folder'], str(output / 'new-pack'))
            self.assertEqual(sentinel.read_text(), 'cloud computing unrelated existing output')
            self.assertFalse((self.source / 'new-pack').exists())
            with self.assertRaises(ValueError):
                kernel.registry.path(str(output.parent / 'ungranted.txt'), result)
        finally:
            kernel.close()

    def test_new_output_folder_export_relative_filename_and_restart(self):
        output = Path(self.temp.name) / 'new-save-folder' / 'reports'
        with patch('services.connections.CredentialStore.get', return_value={}):
            kernel = Runtime(self.config)
        try:
            decision = {'kind': 'tool', 'tool': 'document_export', 'summary': 'Save report', 'complete_after': True,
                'arguments': {'destination': 'brief.md', 'title': 'Actual brief', 'content': 'Cloud meeting notes.'}}
            with patch.object(kernel.planner, 'decide', return_value=decision):
                task = kernel.submit('Save a brief', output_folder=str(output), auto_outputs=True)
                self.assertFalse(output.exists())
                kernel.tick()
                kernel.tick()
            self.assertEqual(kernel.tasks[task['id']]['state'], 'COMPLETED')
            self.assertTrue((output / 'brief.md').is_file())
            self.assertFalse((self.config.workspace / 'brief.md').exists())
            kernel.close()
            with patch('services.connections.CredentialStore.get', return_value={}):
                kernel = Runtime(self.config)
            self.assertEqual(kernel.tasks[task['id']]['output_folder'], str(output))
        finally:
            kernel.close()

    def test_invalid_or_private_output_locations_are_rejected(self):
        with patch('services.connections.CredentialStore.get', return_value={}):
            kernel = Runtime(self.config)
        try:
            for value in ('relative/folder', '', [], str(self.config.root / 'private'), str(self.source / 'holiday.txt')):
                with self.subTest(value=value), self.assertRaises(ValueError):
                    kernel.submit('Save report', output_folder=value)
            self.assertEqual(len(kernel.tasks), 0)
        finally:
            kernel.close()

    def test_boundaries_exclude_credentials_hidden_and_system_directories(self):
        for name in ('credentials.json', '.secret.txt', 'Windows/cloud.txt'):
            path = self.source / name
            path.parent.mkdir(exist_ok=True)
            path.write_text('cloud')
            with self.assertRaises(ValueError):
                self.registry.path(str(path), self.task)
        with self.assertRaises(ValueError):
            self.call('files_collect', query='cloud', destination=str(Path(self.temp.name) / 'outside'))
        self.assertEqual(self.call('files_find_related', query='cloud')['matched'], 2)

    def test_all_export_formats_reopen_and_preserve_sources(self):
        from pypdf import PdfReader
        from docx import Document
        from pptx import Presentation
        from openpyxl import load_workbook
        before = self.manifest()
        for kind in FORMATS:
            with self.subTest(format=kind):
                result = self.call('document_export', destination='reports/summary.' + kind,
                                   title='Cloud computing', content='## Overview\nElastic capacity and cost control.\n- Shared responsibility\n=1+1')
                path = Path(result['artifact'])
                self.assertGreater(path.stat().st_size, 30)
                if kind == 'pdf':
                    self.assertIn('Elastic capacity', ''.join(p.extract_text() for p in PdfReader(path).pages))
                elif kind == 'docx':
                    self.assertIn('Elastic capacity', '\n'.join(p.text for p in Document(path).paragraphs))
                elif kind == 'pptx':
                    self.assertIn('Elastic capacity', '\n'.join(s.text for slide in Presentation(path).slides for s in slide.shapes if s.has_text_frame))
                elif kind == 'xlsx':
                    book = load_workbook(path)
                    self.assertEqual(book.active['B4'].data_type, 's')
                    self.assertTrue(all(c.data_type != 'f' for row in book.active for c in row))
                    book.close()
                elif kind == 'json':
                    self.assertEqual(json.loads(path.read_text())['title'], 'Cloud computing')
                elif kind in ('epub', 'odt', 'zip'):
                    with zipfile.ZipFile(path) as bundle:
                        self.assertIsNone(bundle.testzip())
                        for name in bundle.namelist():
                            if name.endswith(('.xml', '.xhtml', '.opf')):
                                ET.fromstring(bundle.read(name))
                elif kind == 'csv':
                    self.assertIn("'=1+1", path.read_text(encoding='utf-8-sig'))
                elif kind == 'rtf':
                    self.assertTrue(path.read_bytes().startswith(b'{\\rtf1'))
        self.assertEqual(before, self.manifest())

    def test_export_never_overwrites_and_rejects_unsupported_formats(self):
        path = self.source / 'holiday.txt'
        for target in (str(path), 'report.exe', '../outside.pdf'):
            with self.assertRaises(ValueError):
                self.call('document_export', destination=target, title='Title', content='Replacement')
        self.assertEqual(path.read_text(), 'Beach and sunshine.')

    def test_model_title_is_not_duplicated_in_export(self):
        result = self.call('document_export', destination='brief.md', title='Meeting brief', content='# Meeting brief\n\nActual content.')
        self.assertEqual(Path(result['artifact']).read_text().count('Meeting brief'), 1)

    def test_direct_collection_needs_no_model_and_completes_one_quantum(self):
        with patch('services.connections.CredentialStore.get', return_value={}):
            kernel = Runtime(self.config)
        try:
            with patch.object(kernel.planner, 'decide', side_effect=AssertionError('No AI needed')):
                task = kernel.submit('Gather cloud files', auto_outputs=True,
                    workflow={'type': 'collect', 'query': 'cloud', 'destination': 'direct-pack'})
                kernel.tick()
                result = kernel.tasks[task['id']]
                self.assertEqual(result['state'], 'COMPLETED')
                self.assertEqual(result['turns'], 0)
                self.assertEqual(result['steps'][0]['result']['copied'], 2)
            # Opting out of the narrow output grant must still require approval.
            task = kernel.submit('Collect with approval', workflow={'type': 'collect', 'query': 'cloud', 'destination': 'approval-pack'})
            self.assertEqual(task['state'], 'APPROVAL')
            task = kernel.submit('Read permission was removed', auto_outputs=True, scopes=['system'],
                workflow={'type': 'collect', 'query': 'cloud', 'destination': 'ungranted-pack'})
            self.assertEqual(task['state'], 'APPROVAL')
            self.assertFalse(task['pending']['approved'])
        finally:
            kernel.close()

    def test_entire_system_option_expands_only_explicitly_and_pins_provider(self):
        with patch('services.connections.CredentialStore.get', return_value={}):
            kernel = Runtime(self.config)
        try:
            with patch.object(self.config, 'system_roots', return_value=[self.temp.name]):
                local = kernel.submit('Local only')
                self.config.values['provider'] = 'gemini'
                broad = kernel.submit('Whole computer', access_mode='system')
                self.assertNotIn(self.temp.name, local['roots'])
                self.assertIn(self.temp.name, broad['roots'])
                self.assertEqual(local['provider'], 'ollama')
                self.assertEqual(broad['provider'], 'gemini')
        finally:
            kernel.close()

    def test_auto_output_grant_does_not_approve_moves(self):
        with patch('services.connections.CredentialStore.get', return_value={}):
            kernel = Runtime(self.config)
        try:
            move = {'kind': 'tool', 'tool': 'files_move', 'summary': 'Move file',
                    'arguments': {'source': str(self.source / 'holiday.txt'), 'destination': 'moved.txt'}}
            with patch.object(kernel.planner, 'decide', return_value=move):
                task = kernel.submit('Move', auto_outputs=True)
                kernel.tick()
                self.assertEqual(kernel.tasks[task['id']]['state'], 'APPROVAL')
                self.assertTrue((self.source / 'holiday.txt').exists())
        finally:
            kernel.close()

    def test_gemini_protocol_and_key_never_appear_in_public_config(self):
        self.config.values['provider'] = 'gemini'
        router = PlannerRouter(self.config)
        task = {**self.task, 'provider': 'gemini', 'goal': 'Explain cloud', 'history': [], 'steps': []}
        answer = {'kind': 'finish', 'summary': 'Useful answer', 'tool': '', 'arguments': {}}
        with patch.object(router.gemini, 'generate', return_value=(answer, {'promptTokenCount': 12})):
            self.assertEqual(router.decide(task, self.registry.catalog())['metrics']['input_tokens'], 12)
        for invalid in ({'kind': 'shell', 'summary': 'Bad'}, {'kind': 'tool', 'summary': 'Bad', 'arguments': []},
                        {'kind': 'finish', 'summary': 'Bad', 'complete_after': 'true'}):
            with patch.object(router.gemini, 'generate', return_value=(invalid, {})):
                with self.assertRaises(ConnectionError):
                    router.decide(task, self.registry.catalog())
        with patch.object(AIKeys, 'get', return_value='fake-test-key-only'):
            public = self.config.public()
            self.assertTrue(public['gemini_configured'])
            self.assertNotIn('fake-test-key-only', json.dumps(public))
        # Existing local intent never silently migrates to the selected cloud provider.
        with patch.object(router.local, 'decide', return_value=answer) as local, patch.object(router.gemini, 'generate') as cloud:
            router.decide({**task, 'provider': 'ollama'}, [])
            local.assert_called_once()
            cloud.assert_not_called()

    def test_gemini_errors_hide_provider_response_and_use_header_only(self):
        import urllib.error
        from unittest.mock import MagicMock
        gemini = Gemini(self.config)
        opener = MagicMock()
        opener.open.side_effect = urllib.error.HTTPError('https://generativelanguage.googleapis.com', 403,
            'fake-secret-must-not-escape', {}, io.BytesIO(b'fake-secret-must-not-escape'))
        with patch.object(AIKeys, 'get', return_value='fake-secret-must-not-escape'), patch('urllib.request.build_opener', return_value=opener):
            with self.assertRaises(ConnectionError) as error:
                gemini.call('models')
            self.assertNotIn('fake-secret', str(error.exception))
            request = opener.open.call_args.args[0]
            self.assertNotIn('fake-secret', request.full_url)
            self.assertEqual(request.get_header('X-goog-api-key'), 'fake-secret-must-not-escape')


if __name__ == '__main__':
    unittest.main()
