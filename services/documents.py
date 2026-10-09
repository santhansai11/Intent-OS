"""Format-specific exports from plain Markdown. No shell execution or source overwrites."""
import csv
import html
import io
import json
import os
from pathlib import Path
import re
import tempfile
import uuid
import zipfile

FORMATS = ('pdf', 'docx', 'pptx', 'xlsx', 'txt', 'md', 'html', 'csv', 'json', 'rtf', 'odt', 'epub', 'zip')


def blocks(content):
    result = []
    for line in content.splitlines():
        line = line.strip()
        if not line:
            continue
        heading = re.match(r'^(#{1,3})\s+(.+)', line)
        if heading:
            result.append(('heading', heading[2]))
        elif re.match(r'^[-*+]\s+', line):
            result.append(('bullet', line[2:]))
        else:
            result.append(('text', line))
    return result


def plain(value):
    return re.sub(r'\*\*(.*?)\*\*', r'\1', value).replace('`', '')


class DocumentExporter:
    def __init__(self, registry):
        self.registry = registry

    def export(self, args, task):
        target = self.registry.path(args['destination'], task)
        kind = target.suffix.lower().lstrip('.')
        if kind not in FORMATS:
            raise ValueError('Supported formats: ' + ', '.join(FORMATS))
        if target.exists():
            raise ValueError('Choose a new output filename; existing files are preserved')
        title, content = args['title'].strip(), args['content'].strip()
        if not title or len(title) > 200 or not content or len(content) > 60000:
            raise ValueError('Use a title of 1-200 characters and content of 1-60000 characters')
        data = self.render(kind, title, content)
        target.parent.mkdir(parents=True, exist_ok=True)
        fd, temp = tempfile.mkstemp(prefix='.intentos-export-', dir=target.parent)
        try:
            with os.fdopen(fd, 'wb') as file:
                file.write(data)
                file.flush()
                os.fsync(file.fileno())
            os.link(temp, target)
        finally:
            if os.path.exists(temp):
                os.unlink(temp)
        return {'status': 'exported', 'artifact': str(target), 'format': kind, 'title': title,
                'bytes': len(data), 'source_files_modified': False}

    def render(self, kind, title, content):
        lines = content.splitlines()
        first = blocks(lines[0]) if lines else []
        if first and first[0][0] == 'heading' and plain(first[0][1]).casefold() == title.casefold():
            content = '\n'.join(lines[1:]).lstrip()
        rows = blocks(content)
        stream = io.BytesIO()
        if kind == 'md':
            return ('# ' + title + '\n\n' + content + '\n').encode('utf-8')
        if kind == 'txt':
            return (title + '\n\n' + '\n\n'.join(plain(t) for _, t in rows) + '\n').encode('utf-8')
        if kind == 'json':
            return json.dumps({'title': title, 'sections': [{'type': k, 'text': t} for k, t in rows]}, indent=2, ensure_ascii=False).encode('utf-8')
        if kind == 'csv':
            text = io.StringIO(newline='')
            writer = csv.writer(text)
            # Prevent spreadsheet formula execution when CSV is opened.
            def safe(t):
                return "'" + t if t.lstrip().startswith(('=', '+', '-', '@')) else t
            writer.writerow(['type', 'text'])
            writer.writerow(['title', safe(title)])
            writer.writerows((k, safe(plain(t))) for k, t in rows)
            return text.getvalue().encode('utf-8-sig')
        if kind in ('html', 'epub'):
            body = '<h1>' + html.escape(title) + '</h1>'
            for tag, text in rows:
                tag = 'h2' if tag == 'heading' else 'p'
                body += f'<{tag}>' + html.escape(plain(text)) + f'</{tag}>'
            page = ('<?xml version="1.0" encoding="utf-8"?>\n' if kind == 'epub' else '<!doctype html>') + \
                '<html xmlns="http://www.w3.org/1999/xhtml"><head><meta charset="utf-8"/><title>' + html.escape(title) + \
                '</title><style>body{max-width:760px;margin:48px auto;padding:0 24px;font:17px/1.65 sans-serif;color:#183d35}h1{line-height:1.15}h2{margin-top:1.8em}p{overflow-wrap:anywhere}</style></head><body>' + body + '</body></html>'
            if kind == 'html':
                return page.encode('utf-8')
            ident = str(uuid.uuid4())
            with zipfile.ZipFile(stream, 'w') as archive:
                archive.writestr('mimetype', 'application/epub+zip', compress_type=zipfile.ZIP_STORED)
                archive.writestr('META-INF/container.xml', '<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>')
                archive.writestr('content.xhtml', page)
                archive.writestr('nav.xhtml', '<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops"><head><title>Contents</title></head><body><nav epub:type="toc"><ol><li><a href="content.xhtml">' + html.escape(title) + '</a></li></ol></nav></body></html>')
                archive.writestr('content.opf', '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="id"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="id">urn:uuid:' + ident + '</dc:identifier><dc:title>' + html.escape(title) + '</dc:title><dc:language>en</dc:language><meta property="dcterms:modified">' + __import__('datetime').datetime.now(__import__('datetime').timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ') + '</meta></metadata><manifest><item id="main" href="content.xhtml" media-type="application/xhtml+xml"/><item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/></manifest><spine><itemref idref="main"/></spine></package>')
        elif kind == 'rtf':
            def escape(text):
                value = ''
                encoded = text.encode('utf-16-le')
                for unit in [int.from_bytes(encoded[i:i+2], 'little') for i in range(0, len(encoded), 2)]:
                    value += ('\\' + chr(unit)) if unit in (92,123,125) else (chr(unit) if 32 <= unit < 127 else f'\\u{unit if unit < 32768 else unit-65536}?')
                return value
            return ('{\\rtf1\\ansi\\deff0{\\fonttbl{\\f0 Arial;}}\\uc1\\fs24\\b ' + escape(title) + '\\b0\\par ' + '\\par\n'.join(escape(plain(t)) for _, t in rows) + '}').encode('ascii')
        elif kind == 'odt':
            with zipfile.ZipFile(stream, 'w') as archive:
                archive.writestr('mimetype', 'application/vnd.oasis.opendocument.text', compress_type=zipfile.ZIP_STORED)
                paragraphs = '<text:h text:outline-level="1">' + html.escape(title) + '</text:h>'
                paragraphs += ''.join('<text:p>' + html.escape(plain(t)) + '</text:p>' for _, t in rows)
                archive.writestr('content.xml', '<?xml version="1.0" encoding="UTF-8"?><office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" office:version="1.2"><office:body><office:text>' + paragraphs + '</office:text></office:body></office:document-content>')
                archive.writestr('META-INF/manifest.xml', '<?xml version="1.0"?><manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.2"><manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.text"/><manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/></manifest:manifest>')
        elif kind == 'zip':
            with zipfile.ZipFile(stream, 'w', zipfile.ZIP_DEFLATED) as archive:
                archive.writestr('summary.md', '# ' + title + '\n\n' + content)
        elif kind == 'docx':
            from docx import Document
            from docx.shared import Pt
            document = Document()
            document.styles['Normal'].font.name = 'Calibri'
            document.styles['Normal'].font.size = Pt(11)
            document.add_heading(title, 0)
            for k, t in rows:
                if k == 'heading':
                    document.add_heading(plain(t), 1)
                else:
                    document.add_paragraph(plain(t), style='List Bullet' if k == 'bullet' else None)
            document.save(stream)
        elif kind == 'xlsx':
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
            book = Workbook()
            sheet = book.active
            sheet.title = 'Summary'
            sheet.append(['Section', title])
            sheet.freeze_panes = 'A2'
            sheet.column_dimensions['A'].width, sheet.column_dimensions['B'].width = 18, 100
            for k, t in rows:
                # Split long text into readable rows, never formulas.
                for start in range(0, len(t), 800):
                    sheet.append([k, plain(t[start:start+800])])
                    sheet.row_dimensions[sheet.max_row].height = max(28, min(210, 16 * (len(t[start:start+800]) // 95 + 2)))
            for row in sheet:
                for cell in row:
                    cell.data_type = 's'
                    cell.alignment = Alignment(wrap_text=True, vertical='top')
            for cell in sheet[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='183D35')
            book.save(stream)
        elif kind == 'pptx':
            from pptx import Presentation
            from pptx.util import Inches, Pt
            from pptx.dml.color import RGBColor
            import textwrap
            deck = Presentation()
            deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
            current, lines = title, []
            def slide(heading, points):
                s = deck.slides.add_slide(deck.slide_layouts[6])
                s.background.fill.solid()
                s.background.fill.fore_color.rgb = RGBColor.from_string('F7F9F5')
                box = s.shapes.add_textbox(Inches(.7), Inches(.55), Inches(11.9), Inches(1.35)).text_frame
                box.word_wrap = True
                p = box.paragraphs[0]; p.text = heading; p.font.size = Pt(28); p.font.bold = True
                box = s.shapes.add_textbox(Inches(.8), Inches(2.05), Inches(11.7), Inches(4.7)).text_frame
                box.word_wrap = True
                for i, point in enumerate(points):
                    p = box.paragraphs[0] if i == 0 else box.add_paragraph()
                    p.text = point; p.font.size = Pt(20); p.space_after = Pt(14)
            for k, t in rows:
                if k == 'heading':
                    if lines:
                        slide(current, lines); lines = []
                    current = plain(t)[:150]
                else:
                    for part in textwrap.wrap(plain(t), width=150, break_long_words=True):
                        lines.append(part)
                        if len(lines) == 4:
                            slide(current, lines); lines = []
            if lines or not deck.slides:
                slide(current, lines)
            deck.save(stream)
        elif kind == 'pdf':
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.lib import colors
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont
            font = 'Helvetica'
            for path in (Path('C:/Windows/Fonts/arial.ttf'), Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')):
                if path.exists():
                    font = 'IntentSans'
                    if font not in pdfmetrics.getRegisteredFontNames():
                        pdfmetrics.registerFont(TTFont(font, str(path)))
                    break
            styles = getSampleStyleSheet()
            for style in styles.byName.values():
                style.fontName = font
            styles['Title'].textColor = colors.HexColor('#183d35')
            styles['BodyText'].fontSize, styles['BodyText'].leading = 11, 16
            story = [Paragraph(html.escape(title), styles['Title']), Spacer(1, 16)]
            for k, t in rows:
                style = styles['Heading2'] if k == 'heading' else styles['BodyText']
                story.extend([Paragraph(html.escape(plain(t)), style, bulletText='-' if k == 'bullet' else None), Spacer(1, 7)])
            def footer(canvas, document):
                canvas.setFont(font, 9); canvas.setFillColor(colors.HexColor('#677565'))
                canvas.drawString(48, 28, 'IntentOS'); canvas.drawRightString(document.pagesize[0]-48, 28, str(document.page))
            SimpleDocTemplate(stream, rightMargin=48, leftMargin=48, topMargin=48, bottomMargin=48).build(story, onFirstPage=footer, onLaterPages=footer)
        return stream.getvalue()
